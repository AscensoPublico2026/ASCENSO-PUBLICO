# -*- coding: utf-8 -*-
"""Genera el INDICE.md de la biblioteca de manuales cruzando con la matriz."""
import openpyxl, os, glob, re

wb = openpyxl.load_workbook("Convocatorias_Procuraduria.xlsx", data_only=True)
ws = wb["Convocatorias"]
info = {}
for r in range(2, ws.max_row+1):
    cod = ws.cell(r,1).value
    if cod is None: continue
    m = re.match(r"(\d{1,3})", str(cod).strip())
    if not m: continue
    key = m.group(1).zfill(2)+"-2026"
    info[key] = {
        "cargo": ws.cell(r,2).value, "nivel": ws.cell(r,3).value, "grado": ws.cell(r,4).value,
        "plazas": ws.cell(r,5).value, "salario": ws.cell(r,6).value,
        "estudio": (str(ws.cell(r,15).value or "")[:90]).replace("\n"," "),
        "exp": (str(ws.cell(r,16).value or "")[:45]).replace("\n"," "),
    }

DEST = "referencias/manuales-funciones"
tengo = sorted([os.path.basename(f).replace(".pdf","") for f in glob.glob(DEST+"/*.pdf")])

def money(v):
    try: return f"${int(float(v)):,}"
    except: return str(v)

lines = []
lines.append("# 📚 Biblioteca central de Manuales de Funciones — PGN 2026\n")
lines.append("> Manuales de funciones descargados y guardados UNA sola vez. Como las convocatorias se")
lines.append("> repiten entre aspirantes, aquí quedan centralizados para NO volver a subirlos.")
lines.append("> Antes de pedirle un manual a un aspirante, revisa esta lista: si ya está, se reutiliza.\n")
lines.append(f"**Manuales guardados: {len(tengo)}** (de 296 convocatorias totales)\n")
lines.append("| Convocatoria | Cargo | Nivel | Grado | Salario | Plazas | Estudio (resumen) | Exp. |")
lines.append("|---|---|---|---|---|---|---|---|")
for code in sorted(tengo, key=lambda x:int(x.split('-')[0])):
    d = info.get(code, {})
    lines.append(f"| **{code}** | {d.get('cargo','?')} | {d.get('nivel','?')} | {d.get('grado','?')} | {money(d.get('salario','?'))} | {d.get('plazas','?')} | {d.get('estudio','?')} | {d.get('exp','?')} |")

# faltantes (utiles para saber que aun no tenemos)
todos = sorted(info.keys(), key=lambda x:int(x.split('-')[0]))
faltan = [c for c in todos if c not in tengo]
lines.append(f"\n## ⛔ Convocatorias SIN manual aún en la biblioteca ({len(faltan)})")
lines.append("> Cuando un aspirante necesite una de estas, se le pide el manual y luego se agrega aquí con el script de consolidación.\n")
lines.append(", ".join(faltan))

with open(DEST+"/INDICE.md","w",encoding="utf-8") as f:
    f.write("\n".join(lines)+"\n")
print("INDICE.md generado con", len(tengo), "manuales.")
